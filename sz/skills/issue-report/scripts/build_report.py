#!/usr/bin/env python3
"""
SEUZ 총괄 주간현안 1페이지 상세보고서 생성기.

입력: 구조화 JSON (스펙은 references/format-spec.md 참조)
출력: .docx (+ 선택적으로 .md 미리보기)

원본 문서군(W33·W34 최신 표준)의 판형·들여쓰기·주석 텍스트박스를 그대로 재현한다.
BatangChe가 고정폭 폰트라는 점을 이용해 각 본문 줄의 끝 좌표를 계산하고,
그 오른쪽에 파란 주석 텍스트박스를 부유 배치한다.

사용법:
    python3 build_report.py input.json output.docx [--md preview.md]
"""

import argparse
import json
import sys
import unicodedata

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

# ─────────────────────────────────────────────────────────────
# 판형 상수 (최신 표준 = W27·W33·W34)
# ─────────────────────────────────────────────────────────────
EMU_PER_PT = 12700
EMU_PER_TWIP = 635
TWIP_PER_CM = 566.93

PAGE_W, PAGE_H = 11906, 16838          # A4 twips
MARGIN_TOP = MARGIN_BOTTOM = 1418      # 2.5cm
MARGIN_LEFT = MARGIN_RIGHT = 1134      # 2.0cm
HEADER_DIST = FOOTER_DIST = 851        # 1.5cm

TEXT_WIDTH_TWIPS = PAGE_W - MARGIN_LEFT - MARGIN_RIGHT   # 9638 twips ≈ 17.0cm
TEXT_WIDTH_EMU = TEXT_WIDTH_TWIPS * EMU_PER_TWIP

FONT = "BatangChe"
BODY_PT = 14.0
TITLE_PT = 20.0
NOTE_PT = 10.0
NOTE_COLOR = "0000FF"
NOTE_TRACKING = -12          # w:spacing val, 1/20 pt 단위 → -0.6pt
TABLE_PT = 12.0
CAPTION_PT = 12.0
FOOTER_PT = 12.0

TITLE_SPACE_AFTER = 560      # 28pt
BODY_SPACE_AFTER = 180       # 9pt
HEAD_SPACE_AFTER = 180
LINE_RULE = 240              # 100% 줄간격

NOTE_GAP_EMU = 90_000        # 주석끼리 최소 간격 (약 0.25cm)
NOTE_V_OFFSET = 200_000      # 문단 상단 기준 한 줄 아래(약 0.56cm) = 줄간 공백대

# 레벨별 마커와 첫줄 들여쓰기(twips)
LEVELS = {
    "head":  {"marker": "□ ", "first_line": 0,   "bold": True},
    "dash":  {"marker": "- ",  "first_line": 301, "bold": False},
    "dot":   {"marker": "·",   "first_line": 420, "bold": False},
    "arrow": {"marker": "→ ",  "first_line": 560, "bold": False},
    "star":  {"marker": "※ ",  "first_line": 560, "bold": False},
    "tri":   {"marker": "▷ ",  "first_line": 700, "bold": False},
    "cont":  {"marker": "",    "first_line": 560, "bold": False},  # 앞줄 이어쓰기
}

# 1페이지 판정용 (pt)
USABLE_HEIGHT_PT = (PAGE_H - MARGIN_TOP - MARGIN_BOTTOM) / 20.0   # ≈ 700pt
BODY_LINE_PT = 18.0          # BatangChe 14pt 한 줄 높이
TITLE_LINE_PT = 26.0
TABLE_ROW_PT = 20.0


# ─────────────────────────────────────────────────────────────
# 문자 폭 계산 (BatangChe = 고정폭)
# ─────────────────────────────────────────────────────────────
def is_wide(ch: str) -> bool:
    """한글·한자·전각기호는 1em, 라틴/숫자는 0.5em."""
    return unicodedata.east_asian_width(ch) in ("W", "F", "A")


def text_width_pt(text: str, size_pt: float, tracking_pt: float = 0.0) -> float:
    total = 0.0
    for ch in text:
        total += (size_pt if is_wide(ch) else size_pt / 2.0) + tracking_pt
    return total


def line_end_emu(text: str, first_line_twips: int) -> int:
    """본문 한 줄이 끝나는 x좌표(왼쪽 여백 기준 EMU)."""
    return first_line_twips * EMU_PER_TWIP + int(text_width_pt(text, BODY_PT) * EMU_PER_PT)


def anchor_x_emu(text: str, first_line_twips: int, anchor: str) -> int:
    """본문 줄 안에서 anchor 문자열이 시작하는 x좌표(EMU).

    anchor 가 없거나 찾지 못하면 줄 끝을 돌려준다.
    """
    if not anchor:
        return line_end_emu(text, first_line_twips)
    idx = text.find(anchor)
    if idx < 0:
        return line_end_emu(text, first_line_twips)
    return first_line_twips * EMU_PER_TWIP + int(
        text_width_pt(text[:idx], BODY_PT) * EMU_PER_PT)


def wrapped_line_count(text: str, first_line_twips: int) -> int:
    """줄바꿈 없이 몇 줄을 차지하는지."""
    avail = TEXT_WIDTH_EMU - first_line_twips * EMU_PER_TWIP
    used = int(text_width_pt(text, BODY_PT) * EMU_PER_PT)
    return max(1, -(-used // max(avail, 1)))


# ─────────────────────────────────────────────────────────────
# XML 조각
# ─────────────────────────────────────────────────────────────
NS_DECL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
)


def xml_escape(s: str) -> str:
    return (s.replace("&", "&amp;").replace("<", "&lt;")
             .replace(">", "&gt;").replace('"', "&quot;"))


def note_anchor_xml(text: str, x_emu: int, y_emu: int, shape_id: int):
    """본문 오른쪽에 놓이는 파란 주석 텍스트박스."""
    w = int(text_width_pt(text, NOTE_PT, NOTE_TRACKING / 20.0) * EMU_PER_PT) + 220_000
    h = 190_500  # 0.53cm, 10pt 한 줄
    sz = int(NOTE_PT * 2)
    body = xml_escape(text)
    return parse_xml(f'''<w:r {NS_DECL}><w:rPr><w:noProof/></w:rPr><w:drawing>
<wp:anchor distT="45720" distB="45720" distL="114300" distR="114300" simplePos="0"
 relativeHeight="{251_700_000 + shape_id}" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
<wp:simplePos x="0" y="0"/>
<wp:positionH relativeFrom="margin"><wp:posOffset>{x_emu}</wp:posOffset></wp:positionH>
<wp:positionV relativeFrom="paragraph"><wp:posOffset>{y_emu}</wp:posOffset></wp:positionV>
<wp:extent cx="{w}" cy="{h}"/>
<wp:effectExtent l="0" t="0" r="0" b="0"/>
<wp:wrapNone/>
<wp:docPr id="{900 + shape_id}" name="Note{900 + shape_id}"/>
<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="0"/></wp:cNvGraphicFramePr>
<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
<wps:wsp><wps:cNvSpPr txBox="1"><a:spLocks noChangeArrowheads="1"/></wps:cNvSpPr>
<wps:spPr bwMode="auto"><a:xfrm><a:off x="0" y="0"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>
<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/>
<a:ln w="9525"><a:noFill/><a:miter lim="800000"/><a:headEnd/><a:tailEnd/></a:ln></wps:spPr>
<wps:txbx><w:txbxContent><w:p>
<w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/>
<w:rPr><w:rFonts w:ascii="{FONT}" w:eastAsia="{FONT}" w:hAnsi="{FONT}"/>
<w:color w:val="{NOTE_COLOR}"/><w:spacing w:val="{NOTE_TRACKING}"/>
<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr></w:pPr>
<w:r><w:rPr><w:rFonts w:ascii="{FONT}" w:eastAsia="{FONT}" w:hAnsi="{FONT}"/>
<w:color w:val="{NOTE_COLOR}"/><w:spacing w:val="{NOTE_TRACKING}"/>
<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr>
<w:t xml:space="preserve">{body}</w:t></w:r></w:p></w:txbxContent></wps:txbx>
<wps:bodyPr rot="0" vert="horz" wrap="none" lIns="0" tIns="0" rIns="0" bIns="0"
 anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr></wps:wsp>
</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>''')


# ─────────────────────────────────────────────────────────────
# 문서 조립
# ─────────────────────────────────────────────────────────────
def setup_document(doc: Document, footer_text: str = "- 1 -"):
    sec = doc.sections[0]
    sectPr = sec._sectPr
    pgSz = sectPr.find(qn("w:pgSz"))
    pgSz.set(qn("w:w"), str(PAGE_W))
    pgSz.set(qn("w:h"), str(PAGE_H))
    pgMar = sectPr.find(qn("w:pgMar"))
    for k, v in (("top", MARGIN_TOP), ("bottom", MARGIN_BOTTOM),
                 ("left", MARGIN_LEFT), ("right", MARGIN_RIGHT),
                 ("header", HEADER_DIST), ("footer", FOOTER_DIST), ("gutter", 0)):
        pgMar.set(qn("w:" + k), str(v))

    # 문서 기본 서식: BatangChe 14pt
    styles = doc.styles.element
    dd = styles.find(qn("w:docDefaults"))
    rpr_default = dd.find(qn("w:rPrDefault"))
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = parse_xml(f'<w:rPr xmlns:w="{qn("w:x")[1:-2]}"/>')
    rf_def = rpr.find(qn("w:rFonts"))
    if rf_def is not None:
        for a in ("w:asciiTheme", "w:eastAsiaTheme", "w:hAnsiTheme"):
            if rf_def.get(qn(a)) is not None:
                del rf_def.attrib[qn(a)]
    for tag, attrs in (("w:rFonts", {"w:ascii": FONT, "w:eastAsia": FONT, "w:hAnsi": FONT}),
                       ("w:color", {"w:val": "000000"}),
                       ("w:kern", {"w:val": "2"}),
                       ("w:sz", {"w:val": str(int(BODY_PT * 2))}),
                       ("w:szCs", {"w:val": str(int(BODY_PT * 2))})):
        el = rpr.find(qn(tag))
        if el is None:
            el = parse_xml(f'<{tag} xmlns:w="http://schemas.openxmlformats.org/'
                           f'wordprocessingml/2006/main"/>')
            rpr.append(el)
        for k, v in attrs.items():
            el.set(qn(k), v)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    rPr = normal.element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    for a in ("w:ascii", "w:eastAsia", "w:hAnsi"):
        rf.set(qn(a), FONT)

    # 꼬리말
    footer = sec.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    set_ppr(p, jc="center", after=0, line=LINE_RULE)
    add_run(p, footer_text, size_pt=FOOTER_PT)


def set_ppr(p, *, jc=None, first_line=None, after=None, before=None,
            line=LINE_RULE, left=None):
    pPr = p._p.get_or_add_pPr()
    if jc:
        el = pPr.find(qn("w:jc"))
        if el is None:
            el = parse_xml('<w:jc xmlns:w="http://schemas.openxmlformats.org/'
                           'wordprocessingml/2006/main"/>')
            pPr.append(el)
        el.set(qn("w:val"), jc)
    sp = parse_xml('<w:spacing xmlns:w="http://schemas.openxmlformats.org/'
                   'wordprocessingml/2006/main"/>')
    if after is not None:
        sp.set(qn("w:after"), str(after))
    if before is not None:
        sp.set(qn("w:before"), str(before))
    sp.set(qn("w:line"), str(line))
    sp.set(qn("w:lineRule"), "auto")
    old = pPr.find(qn("w:spacing"))
    if old is not None:
        pPr.remove(old)
    pPr.append(sp)
    if first_line or left:
        ind = parse_xml('<w:ind xmlns:w="http://schemas.openxmlformats.org/'
                        'wordprocessingml/2006/main"/>')
        if first_line:
            ind.set(qn("w:firstLine"), str(first_line))
        if left:
            ind.set(qn("w:left"), str(left))
        pPr.append(ind)


def add_run(p, text, *, size_pt=BODY_PT, bold=False, underline=False, color=None):
    r = p.add_run(text)
    r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    for a in ("w:ascii", "w:eastAsia", "w:hAnsi"):
        rf.set(qn(a), FONT)
    r.font.size = None
    sz = parse_xml(f'<w:sz xmlns:w="http://schemas.openxmlformats.org/'
                   f'wordprocessingml/2006/main" w:val="{int(size_pt * 2)}"/>')
    szcs = parse_xml(f'<w:szCs xmlns:w="http://schemas.openxmlformats.org/'
                     f'wordprocessingml/2006/main" w:val="{int(size_pt * 2)}"/>')
    rPr.append(sz)
    rPr.append(szcs)
    r.bold = bold
    r.underline = underline
    if color:
        r.font.color.rgb = None
        c = parse_xml(f'<w:color xmlns:w="http://schemas.openxmlformats.org/'
                      f'wordprocessingml/2006/main" w:val="{color}"/>')
        rPr.append(c)
    return r


def add_table(doc, block):
    rows = block["rows"]
    caption = block.get("caption")
    if caption:
        p = doc.add_paragraph()
        set_ppr(p, jc="center", after=120)
        add_run(p, caption, size_pt=CAPTION_PT)

    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    # "grid" = 전체 실선(W18·W20형), "minimal" = 헤더 하단선만(W33형)
    if block.get("borders", "grid") == "grid":
        t.style = "Table Grid"
    tblPr = t._tbl.tblPr
    ind = parse_xml('<w:tblInd xmlns:w="http://schemas.openxmlformats.org/'
                    'wordprocessingml/2006/main" w:w="560" w:type="dxa"/>')
    tblPr.append(ind)

    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = t.cell(ri, ci)
            cell.vertical_alignment = 1  # center
            para = cell.paragraphs[0]
            para.text = ""
            set_ppr(para, jc="center", after=0)
            txt = row[ci] if ci < len(row) else ""
            add_run(para, str(txt), size_pt=TABLE_PT, bold=(ri == 0))
        if ri == 0:  # 헤더행 하단 굵은 선
            for ci in range(ncols):
                tcPr = t.cell(0, ci)._tc.get_or_add_tcPr()
                tcPr.append(parse_xml(
                    '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/'
                    'wordprocessingml/2006/main">'
                    '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="auto"/>'
                    '</w:tcBorders>'))

    after = doc.add_paragraph()
    set_ppr(after, after=BODY_SPACE_AFTER)
    return t


def normalize_notes(block):
    """notes / note+anchor 두 표기를 하나로 정리."""
    notes = block.get("notes")
    if notes is None and block.get("note"):
        notes = [{"anchor": block.get("anchor", ""), "text": block["note"]}]
    out = []
    for n in (notes or []):
        out.append({"anchor": "", "text": n} if isinstance(n, str) else n)
    return out


def attach_notes(p, notes, base_x_fn, note_id, warnings, label=""):
    """문단 p 아래에 주석 텍스트박스를 배치한다.

    base_x_fn(anchor) -> EMU. 본문은 들여쓰기 기준, 제목은 가운데정렬 기준으로
    서로 다른 계산식을 넘긴다.
    """
    prev_right = -1
    for nspec in notes:
        note_id += 1
        raw = nspec["text"]
        note_text = raw if raw.lstrip().startswith("*") else "* " + raw
        x = base_x_fn(nspec.get("anchor", ""))
        x += int(nspec.get("dx_cm", 0) * 360000)
        nw = int(text_width_pt(note_text, NOTE_PT, NOTE_TRACKING / 20.0) * EMU_PER_PT)

        if x <= prev_right:                      # 같은 줄 주석끼리 충돌
            x = prev_right + NOTE_GAP_EMU
        if x + nw > TEXT_WIDTH_EMU:              # 우측 여백 초과
            x = max(0, TEXT_WIDTH_EMU - nw)
            warnings.append(f"[주석넘침] '{note_text[:24]}…' 가 우측 여백을 넘어 "
                            f"왼쪽으로 당겼습니다. 주석을 짧게 하거나 앵커를 앞쪽 단어로 옮기세요.")
        prev_right = x + nw
        p._p.append(note_anchor_xml(note_text, x,
                                    NOTE_V_OFFSET + int(nspec.get("dy_cm", 0) * 360000),
                                    note_id))
    return note_id


def title_anchor_x_emu(title: str, anchor: str) -> int:
    """가운데정렬된 제목 안에서 anchor 가 시작하는 x좌표(EMU).

    anchor 가 없으면 제목 시작점(왼쪽 끝)에 맞춘다. 원본 W36이
    제목 아래 좌측에 `* Atradius社` 를 두는 형태를 기본값으로 삼는다.
    """
    tw = int(text_width_pt(title, TITLE_PT) * EMU_PER_PT)
    start = max(0, (TEXT_WIDTH_EMU - tw) // 2)
    if not anchor:
        return start
    idx = title.find(anchor)
    if idx < 0:
        return start
    return start + int(text_width_pt(title[:idx], TITLE_PT) * EMU_PER_PT)


def build(spec, out_path, md_path=None):
    doc = Document()
    setup_document(doc, spec.get("footer", "- 1 -"))

    # 본문 첫 문단(python-docx 기본 빈 문단) 제거
    body = doc.element.body
    for p in body.findall(qn("w:p")):
        body.remove(p)

    warnings = []
    note_id = 0

    # 제목
    title = spec["title"]
    tp = doc.add_paragraph()
    set_ppr(tp, jc="center", after=spec.get("title_space_after", TITLE_SPACE_AFTER))
    add_run(tp, title, size_pt=TITLE_PT, bold=True, underline=True)

    title_notes = normalize_notes({"notes": spec.get("title_notes")})
    for n in title_notes:
        a = n.get("anchor", "")
        if a and a not in title:
            warnings.append(f"[앵커없음] 제목에서 '{a}' 를 찾지 못해 제목 시작점에 배치했습니다.")
    note_id = attach_notes(tp, title_notes,
                           lambda a: title_anchor_x_emu(title, a),
                           note_id, warnings)

    used_pt = TITLE_LINE_PT + spec.get("title_space_after", TITLE_SPACE_AFTER) / 20.0

    for block in spec["blocks"]:
        if block.get("type") == "table":
            add_table(doc, block)
            used_pt += TABLE_ROW_PT * len(block["rows"]) + 30
            continue

        level = block.get("level", "dash")
        cfg = LEVELS.get(level)
        if cfg is None:
            raise ValueError(f"알 수 없는 level: {level}")

        text = cfg["marker"] + block["text"]
        p = doc.add_paragraph()
        set_ppr(p, first_line=cfg["first_line"] or None,
                after=block.get("space_after", HEAD_SPACE_AFTER if level == "head"
                                else BODY_SPACE_AFTER))
        add_run(p, text, size_pt=BODY_PT, bold=cfg["bold"])

        nlines = wrapped_line_count(text, cfg["first_line"])
        if nlines > 1:
            warnings.append(f"[줄넘침] '{text[:24]}…' 가 {nlines}줄로 흐릅니다. "
                            f"본문은 한 줄에 담기도록 압축하세요.")
        used_pt += BODY_LINE_PT * nlines + block.get(
            "space_after", HEAD_SPACE_AFTER if level == "head" else BODY_SPACE_AFTER) / 20.0

        # 주석: 설명 대상 단어 바로 아래 줄간 공백에 배치
        notes = normalize_notes(block)
        for nspec in notes:
            anchor = nspec.get("anchor", "")
            if anchor and anchor not in text:
                warnings.append(f"[앵커없음] '{anchor}' 를 본문 '{block['text'][:20]}…' "
                                f"에서 찾지 못해 줄 끝에 배치했습니다.")
        note_id = attach_notes(p, notes,
                               lambda a: anchor_x_emu(text, cfg["first_line"], a),
                               note_id, warnings)

    if used_pt > USABLE_HEIGHT_PT:
        over = used_pt - USABLE_HEIGHT_PT
        warnings.append(f"[분량초과] 예상 높이 {used_pt:.0f}pt / 가용 {USABLE_HEIGHT_PT:.0f}pt "
                        f"— 약 {over / (BODY_LINE_PT + BODY_SPACE_AFTER / 20):.0f}줄 초과. 본문을 줄이세요.")
    elif used_pt > USABLE_HEIGHT_PT * 0.93:
        warnings.append(f"[분량주의] 예상 높이 {used_pt:.0f}pt / 가용 {USABLE_HEIGHT_PT:.0f}pt — 여유 없음.")

    doc.save(out_path)

    if md_path:
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(render_markdown(spec))

    return warnings


def render_markdown(spec) -> str:
    out = [f"# {spec['title']}", ""]
    for n in normalize_notes({"notes": spec.get("title_notes")}):
        raw = n["text"]
        out.append("      " + (raw if raw.lstrip().startswith("*") else "* " + raw))
    if spec.get("title_notes"):
        out.append("")
    for b in spec["blocks"]:
        if b.get("type") == "table":
            if b.get("caption"):
                out.append(f"**{b['caption']}**")
                out.append("")
            rows = b["rows"]
            out.append("| " + " | ".join(str(c) for c in rows[0]) + " |")
            out.append("|" + "---|" * len(rows[0]))
            for r in rows[1:]:
                out.append("| " + " | ".join(str(c) for c in r) + " |")
            out.append("")
            continue
        lvl = b.get("level", "dash")
        indent = {"head": "", "dash": "  ", "dot": "    ",
                  "arrow": "      ", "star": "      ", "tri": "        ",
                  "cont": "      "}[lvl]
        line = indent + LEVELS[lvl]["marker"] + b["text"]
        notes = b.get("notes")
        if notes is None and b.get("note"):
            notes = [{"anchor": b.get("anchor", ""), "text": b["note"]}]
        for n in (notes or []):
            if isinstance(n, str):
                n = {"anchor": "", "text": n}
            t = n["text"] if n["text"].lstrip().startswith("*") else "* " + n["text"]
            tag = f"[{n['anchor']}]" if n.get("anchor") else ""
            line += f"  `{tag}{t}`"
        out.append(line)
    return "\n".join(out) + "\n"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--md", dest="md", default=None, help="마크다운 미리보기 경로")
    args = ap.parse_args()

    with open(args.input, encoding="utf-8") as f:
        spec = json.load(f)

    warnings = build(spec, args.output, args.md)
    print(f"생성 완료: {args.output}")
    if args.md:
        print(f"미리보기: {args.md}")
    if warnings:
        print("\n── 점검 사항 ──")
        for w in warnings:
            print(" " + w)
    else:
        print("점검 사항 없음 (1페이지 이내, 주석 겹침 없음)")


if __name__ == "__main__":
    main()
